"""Benchmark: this system vs other converters on identical documents.

Usage:  python benchmarks/run_benchmark.py [--extra DIR_WITH_PDFS] [--out benchmarks/results]
Everything reported is measured by this script on the machine it runs on; nothing is hard-coded.

Metrics (all computed from the produced .docx, never from the converter's own claims):
  * text WER / CER  - jiwer on whitespace-normalised text vs ground truth (native: the text used to generate the PDF; scanned: the
                      text of the vector source page). Reading order is part of the score.
  * table F1        - cell-level exact-match F1 between ground-truth cells and the cells of the first table in the .docx.
  * layout fidelity - (Microsoft Word + pywin32 only) the .docx is rendered back to PDF by Word and every word of the original is
                      matched to the same word in the render; reported: page-count ratio, mean position error as % of page width,
                      and the share of words within 2 % of the page width of their original position.
  * time / memory   - wall-clock seconds and peak resident set size (psutil, sampled every 20 ms) of a fresh worker process.
"""
from __future__ import annotations

import argparse
import json
import os
import platform
import subprocess
import sys
import tempfile
import threading
import time

import psutil

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import docx                         # noqa: E402
import jiwer                        # noqa: E402
import pymupdf as fitz              # noqa: E402
import testdata                     # noqa: E402

TOOL_TIMEOUT_S = 180
WORKER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_worker.py")


# --------------------------------------------------------------------------- running a converter in a measured subprocess
def run_tool(tool: str, pdf: str, out: str):
    t0 = time.perf_counter()
    p = subprocess.Popen([sys.executable, "-W", "ignore", WORKER, tool, pdf, out], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    ps = psutil.Process(p.pid)
    peak = 0

    def sample():
        nonlocal peak
        while p.poll() is None:
            try:
                rss = ps.memory_info().rss + sum(c.memory_info().rss for c in ps.children(recursive=True))
                peak = max(peak, rss)
            except psutil.Error:
                pass
            time.sleep(0.02)
    th = threading.Thread(target=sample)
    th.start()
    try:
        o, e = p.communicate(timeout=TOOL_TIMEOUT_S)
        timed_out = False
    except subprocess.TimeoutExpired:
        for c in ps.children(recursive=True) + [ps]:
            try:
                c.kill()
            except psutil.Error:
                pass
        o, e = p.communicate()
        timed_out = True
        if tool == "word_reflow":
            subprocess.run(["taskkill", "/F", "/IM", "WINWORD.EXE"], capture_output=True)
    th.join()
    return {"ok": (not timed_out) and p.returncode == 0 and os.path.exists(out), "timed_out": timed_out, "seconds": round(time.perf_counter() - t0, 2), "peak_mb": round(peak / 1e6, 1),
            "error": (e.decode(errors="ignore")[-300:] if p.returncode else "")}


# --------------------------------------------------------------------------- metrics
def norm(s: str) -> str:
    return " ".join(s.split())


def docx_text(path: str) -> str:
    d = docx.Document(path)
    body = d.element.body
    parts = []
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for ch in body.iterchildren():
        if ch.tag.endswith("}p"):
            parts.append(Paragraph(ch, d).text)
        elif ch.tag.endswith("}tbl"):
            for r in Table(ch, d).rows:
                parts.append(" ".join(c.text for c in r.cells))
    return "\n".join(parts)


def text_scores(truth: str, got: str):
    t, g = norm(truth), norm(got)
    if not t:
        return None
    return {"wer": round(jiwer.wer(t, g), 4), "cer": round(jiwer.cer(t, g), 4)}


def table_f1(truth_rows, path: str):
    d = docx.Document(path)
    if not d.tables:
        return {"f1": 0.0, "found": 0}
    got = [[c.text.strip() for c in r.cells] for r in d.tables[0].rows]
    T = {(i, j, c) for i, r in enumerate(truth_rows) for j, c in enumerate(r)}
    G = {(i, j, c) for i, r in enumerate(got) for j, c in enumerate(r)}
    tp = len(T & G)
    p, r = tp / max(1, len(G)), tp / max(1, len(T))
    return {"f1": round(2 * p * r / (p + r), 4) if p + r else 0.0, "found": len(d.tables)}


def word_render(docx_path: str, pdf_out: str) -> bool:
    try:
        import win32com.client                                     # type: ignore
    except Exception:
        return False
    try:
        w = win32com.client.DispatchEx("Word.Application")
        w.Visible = False
        d = w.Documents.Open(os.path.abspath(docx_path), False, True)
        d.ExportAsFixedFormat(os.path.abspath(pdf_out), 17)
        d.Close(False)
        w.Quit()
        return True
    except Exception:
        return False


def layout_fidelity(orig_pdf: str, rendered_pdf: str):
    a, b = fitz.open(orig_pdf), fitz.open(rendered_pdf)
    errs, near, total = [], 0, 0
    for i in range(min(a.page_count, b.page_count)):
        wa, wb = a[i].get_text("words"), b[i].get_text("words")
        W = a[i].rect.width
        pool = {}
        for w in wb:
            pool.setdefault(w[4], []).append(w)
        for w in wa:
            cands = pool.get(w[4])
            if not cands:
                continue
            best = min(cands, key=lambda c: abs(c[0] - w[0]) + abs(c[1] - w[1]))
            d = ((best[0] - w[0]) ** 2 + (best[1] - w[1]) ** 2) ** 0.5 / W
            errs.append(d); total += 1; near += d <= 0.02
    if not errs:
        return None
    return {"page_ratio": round(b.page_count / a.page_count, 3), "mean_err_pct_width": round(100 * sum(errs) / len(errs), 2),
            "within_2pct": round(near / max(1, sum(len(x.get_text("words")) for x in a)), 3), "matched_words": total}


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "results"))
    ap.add_argument("--extra", default=None, help="folder of your own PDFs (reported without ground truth: time, memory, layout only)")
    ap.add_argument("--tools", default="ours,pdf2docx", help="comma list; word_reflow (Microsoft Word PDF reflow) is optional and can hang on some systems")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    work = tempfile.mkdtemp(prefix="p2d_bench_")
    meta = testdata.build_all(os.path.join(work, "data"))
    docs = {k: os.path.join(work, "data", k + ".pdf") for k in meta}
    truth = {k: meta[k].get("text") for k in meta}
    if a.extra:
        for f in sorted(os.listdir(a.extra)):
            if f.lower().endswith(".pdf"):
                docs["extra:" + f] = os.path.join(a.extra, f); truth["extra:" + f] = None
    tools = a.tools.split(",")
    rows = []
    for name, pdf in docs.items():
        for tool in tools:
            out = os.path.join(work, "%s__%s.docx" % (name.replace(":", "_").replace(" ", "_"), tool))
            r = {"doc": name, "tool": tool, "pages": fitz.open(pdf).page_count}
            r.update(run_tool(tool, pdf, out))
            if r["ok"]:
                if truth.get(name):
                    r["text"] = text_scores(truth[name], docx_text(out))
                if name == "tables":
                    r["table"] = table_f1(meta["tables"]["table"], out)
                rp = out.replace(".docx", "_render.pdf")
                if word_render(out, rp) and os.path.exists(rp):
                    r["layout"] = layout_fidelity(pdf, rp)
            rows.append(r)
            print(name, tool, {k: v for k, v in r.items() if k not in ("error",)}, flush=True)
    env = {"python": sys.version.split()[0], "platform": platform.platform(), "cpu": platform.processor(), "cores": psutil.cpu_count(logical=True),
           "ram_gb": round(psutil.virtual_memory().total / 1e9, 1), "pymupdf": fitz.VersionBind,
           "docx": docx.__version__ if hasattr(docx, "__version__") else "python-docx", "jiwer": jiwer.__version__ if hasattr(jiwer, "__version__") else "?"}
    try:
        import pdf2docx, importlib.metadata as md
        env["pdf2docx"] = md.version("pdf2docx")
    except Exception:
        pass
    try:
        from engine import ocr
        import pytesseract
        env["tesseract"] = str(pytesseract.get_tesseract_version()) if ocr.ocr_available() else "unavailable"
    except Exception:
        env["tesseract"] = "unavailable"
    with open(os.path.join(a.out, "results.json"), "w", encoding="utf-8") as f:
        json.dump({"environment": env, "results": rows}, f, indent=2)
    from plot_results import make_plots
    make_plots(os.path.join(a.out, "results.json"), a.out)
    print("written to", a.out)


if __name__ == "__main__":
    main()
